"""Xem và xuất lịch thi theo docs/data_schema.md mục 7.

Chạy: streamlit run app/app.py
"""
import json
from io import BytesIO
from pathlib import Path

import streamlit as st
from openpyxl import Workbook

DEFAULT_PATH = Path(__file__).resolve().parents[1] / "output" / "toy" / "schedule.json"
PERIODS = {
    1: "07:00 - 09:30",
    2: "09:40 - 12:10",
    3: "13:00 - 15:30",
    4: "15:40 - 18:10",
}
STATUS_NOTES = {
    "OPTIMAL": "Đã tìm được lịch tối ưu theo hàm mục tiêu.",
    "FEASIBLE": "Đã có lịch hợp lệ, nhưng chưa chứng minh được tối ưu trong thời gian cho phép.",
    "INPUT_ERROR": "Dữ liệu đầu vào có lỗi; solver chưa chạy. Hãy sửa các lỗi được liệt kê rồi chạy lại.",
    "INFEASIBLE": "Không có lịch thỏa tất cả ràng buộc cứng. Hãy kiểm tra số phòng, ca trống và giám thị.",
    "UNKNOWN": "Hết thời gian nhưng chưa tìm được lịch hợp lệ. Có thể tăng thời gian giải rồi chạy lại.",
}
PENALTY_NOTES = {
    "A-S1": "Hạn chế sinh viên thi từ hai môn trong một ngày.",
    "A-S2": "Giãn cách các môn thi của cùng sinh viên.",
    "A-S3": "Hạn chế chia một môn ra nhiều phòng.",
    "A-S4": "Hạn chế chỗ trống trong phòng thi.",
    "A-S5": "Chia đều tải coi thi giữa các giám thị.",
    "A-S6": "Hạn chế số lượt thi, tức số đề phải chuẩn bị.",
}
KPI_NOTES = {
    "hard_violations": "Số vi phạm ràng buộc cứng; lịch đúng phải bằng 0.",
    "students_multi_exam_day": "Số sinh viên có ít nhất một ngày bắt đầu từ hai lượt thi.",
    "avg_gap_days": "Khoảng cách ngày giữa các lượt liên tiếp: tính trung bình từng sinh viên rồi trung bình đều trên các sinh viên thi từ hai lượt.",
    "min_gap_days": "Khoảng cách ngày nhỏ nhất giữa hai lượt liên tiếp; cùng ngày tính là 0.",
    "proctor_load_min": "Số ca coi thi ít nhất trên toàn bộ giảng viên; lượt hai ca tính hai ca.",
    "proctor_load_max": "Số ca coi thi nhiều nhất trên toàn bộ giảng viên; lượt hai ca tính hai ca.",
    "proctor_load_range": "Chênh lệch giữa số ca coi thi nhiều nhất và ít nhất.",
    "num_sittings": "Tổng số lượt thi, cũng là số đề cần chuẩn bị.",
    "room_utilization_normal": "Tỷ lệ cặp (phòng thường, ca) được dùng trên các cặp không bận.",
    "room_utilization_lab": "Tỷ lệ cặp (phòng máy, ca) được dùng trên các cặp không bận.",
    "seat_fill_rate": "Tổng (số sinh viên × số ca) chia tổng (sức chứa phòng dùng × số ca).",
    "solve_time_sec": "Thời gian giải tính bằng giây, không gồm thời gian đọc dữ liệu.",
}


def format_value(value, ratio=False):
    if value is None:
        return "Không áp dụng / chưa có"
    return f"{value:.2%}" if ratio else str(value)


def calendar_rows(sittings):
    rows = []
    for date in sorted({s["date"] for s in sittings}):
        row = {"Ngày": date}
        for period, hours in PERIODS.items():
            entries = [
                f"{s['course_id']} - {s['course_name']} | Phòng: "
                + ", ".join(r["room_id"] for r in s["rooms"])
                for s in sittings if s["date"] == date and period in s["periods"]
            ]
            row[f"Ca {period} ({hours})"] = "\n".join(entries) or "—"
        rows.append(row)
    return rows


def room_rows(sittings):
    return [
        {
            "Lượt thi": s["sitting_id"],
            "Mã môn": s["course_id"],
            "Môn": s["course_name"],
            "Ngày": s["date"],
            "Ca": ", ".join(map(str, s["periods"])),
            "Giờ": f"{s['start_time']} - {s['end_time']}",
            "Phòng": r["room_id"],
            "Số SV": r["num_students"],
            "Giám thị": ", ".join(r["proctors"]),
        }
        for s in sittings for r in s["rooms"]
    ]


def export_excel(sittings):
    workbook = Workbook()
    workbook.remove(workbook.active)
    summary = room_rows(sittings)
    students = [
        {**row, "Mã sinh viên": student}
        for row, room in zip(summary, (r for s in sittings for r in s["rooms"]))
        for student in room["students"]
    ]
    headers = ["Lượt thi", "Mã môn", "Môn", "Ngày", "Ca", "Giờ", "Phòng", "Số SV", "Giám thị"]
    for name, rows, columns in (
        ("Lịch tổng", summary, headers),
        ("Sinh viên theo phòng", students, headers + ["Mã sinh viên"]),
    ):
        sheet = workbook.create_sheet(name)
        sheet.append(columns)
        for row in rows:
            sheet.append([row[column] for column in columns])
        # Giữ mã có số 0 đầu và không diễn giải dữ liệu JSON thành công thức Excel.
        for cells in sheet.iter_rows():
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(
                55, max(len(str(cell.value or "")) for cell in column) + 2
            )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def show_messages(schedule):
    st.subheader("Thông báo và vi phạm")
    for level, label, display in (
        ("error", "Lỗi", st.error),
        ("warning", "Cảnh báo", st.warning),
        ("info", "Thông tin", st.info),
    ):
        messages = [m for m in schedule.get("messages", []) if m.get("level") == level]
        if messages:
            st.markdown(f"**{label} ({len(messages)})**")
        for message in messages:
            location = message.get("file") or ""
            if message.get("line") is not None:
                location += f", dòng {message['line']}"
            display(f"[{message.get('code', level)}] {message.get('message', '')}"
                    + (f" ({location})" if location else ""))
    violations = schedule.get("violations", [])
    if violations:
        st.error(f"PHÁT HIỆN {len(violations)} VI PHẠM RÀNG BUỘC CỨNG — không sử dụng lịch trước khi sửa lỗi.")
        for violation in violations:
            st.error(f"[{violation.get('code', '')}] {violation.get('message', '')}")
            if violation.get("refs"):
                st.json(violation["refs"])
    elif not schedule.get("messages"):
        st.info("Không có thông báo hoặc vi phạm được ghi trong file.")


def main():
    st.set_page_config(page_title="Lịch thi", layout="wide")
    st.title("Lịch thi cuối kỳ")
    uploaded = st.file_uploader("Tải file kết quả schedule.json", type=["json"])
    path = st.text_input("Đường dẫn file kết quả", str(DEFAULT_PATH))
    st.caption("File tải lên được ưu tiên; nếu chưa tải, mở đường dẫn trên (mặc định là bộ toy).")
    try:
        content = uploaded.getvalue() if uploaded is not None else Path(path).read_bytes()
        schedule = json.loads(content.decode("utf-8-sig"))
        if not isinstance(schedule, dict) or schedule.get("status") not in STATUS_NOTES:
            raise ValueError("Thiếu trạng thái hợp lệ theo schema mục 7.")
        if not isinstance(schedule.get("messages", []), list) or not isinstance(schedule.get("violations", []), list):
            raise ValueError("messages và violations phải là mảng.")
        if any(not isinstance(item, dict) for key in ("messages", "violations") for item in schedule.get(key, [])):
            raise ValueError("Mỗi thông báo và vi phạm phải là một đối tượng.")
    except (OSError, UnicodeError, ValueError) as exc:
        st.error(f"Không đọc được file kết quả: {exc}")
        return

    status = schedule["status"]
    cols = st.columns(4)
    cols[0].metric("Trạng thái", status, help=STATUS_NOTES[status])
    cols[1].metric("Giá trị mục tiêu", format_value(schedule.get("objective")), help="Tổng điểm phạt đã nhân trọng số; càng nhỏ càng tốt.")
    cols[2].metric("Cận dưới tốt nhất", format_value(schedule.get("best_bound")), help="Cận dưới solver chứng minh được; bằng mục tiêu khi tối ưu.")
    cols[3].metric("Thời gian giải (giây)", format_value(schedule.get("solve_time_sec")), help=KPI_NOTES["solve_time_sec"])
    (st.success if status == "OPTIMAL" else st.info if status == "FEASIBLE" else st.error)(STATUS_NOTES[status])
    show_messages(schedule)
    if status not in ("OPTIMAL", "FEASIBLE"):
        return

    st.subheader("Điểm phạt A-S1 đến A-S6")
    penalties = schedule.get("penalties") or {}
    st.caption("Điểm phạt chưa nhân trọng số; không phải phần đóng góp trực tiếp vào giá trị mục tiêu.")
    st.dataframe([{"Chỉ số": key, "Giá trị": format_value(penalties.get(key)), "Ghi chú": note}
                  for key, note in PENALTY_NOTES.items()], hide_index=True, width="stretch")
    st.subheader("Toàn bộ KPI")
    kpi = schedule.get("kpi") or {}
    st.dataframe([{"Chỉ số": key, "Giá trị": format_value(kpi.get(key), key in ("room_utilization_normal", "room_utilization_lab", "seat_fill_rate")), "Ghi chú": note}
                  for key, note in KPI_NOTES.items()], hide_index=True, width="stretch")
    st.caption("Không áp dụng / chưa có: giá trị null hoặc thiếu trong file; không được coi là 0.")

    sittings = schedule.get("sittings") or []
    if not sittings:
        st.info("File không có lượt thi để hiển thị.")
        return
    sittings = sorted(sittings, key=lambda s: (s["date"], s["start_period"], s["course_id"]))
    rows = room_rows(sittings)
    calendar, rooms, students, proctors = st.tabs(["Lịch ngày × ca", "Theo phòng", "Sinh viên", "Giám thị"])
    with calendar:
        st.subheader("Lịch ngày × ca")
        st.caption("Lượt thi hai ca xuất hiện ở cả hai ô. Giờ hiển thị là khung giữ phòng, không phải giờ nộp bài.")
        st.dataframe(calendar_rows(sittings), hide_index=True, width="stretch", row_height=140)
    with rooms:
        room_id = st.selectbox("Chọn phòng", sorted({row["Phòng"] for row in rows}))
        if room_id is not None:
            st.dataframe([row for row in rows if row["Phòng"] == room_id], hide_index=True, width="stretch")
    with students:
        student_id = st.text_input("Mã sinh viên").strip()
        if student_id:
            matches = [row for row, room in zip(rows, (r for s in sittings for r in s["rooms"])) if student_id in room["students"]]
            if matches:
                st.dataframe(matches, hide_index=True, width="stretch")
            else:
                st.info("Không tìm thấy lịch thi của sinh viên này trong file.")
    with proctors:
        lecturer = st.selectbox("Chọn giảng viên", sorted({g for s in sittings for r in s["rooms"] for g in r["proctors"]}))
        if lecturer is not None:
            matches = [row for row, room in zip(rows, (r for s in sittings for r in s["rooms"])) if lecturer in room["proctors"]]
            st.dataframe(matches, hide_index=True, width="stretch")
        st.caption("File kết quả chỉ có mã giảng viên được phân công, không có tên hoặc danh sách giảng viên không được phân công.")
    st.download_button("Xuất Excel", data=export_excel(sittings), file_name="lich_thi.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


if __name__ == "__main__":
    main()
