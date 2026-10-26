"""Kiểm thử giao diện trực tiếp từ file kết quả, không chạy solver."""
import json
import shutil
import subprocess
import sys
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from app.app import DEFAULT_PATH, KPI_NOTES, calendar_rows, export_excel

ROOT = Path(__file__).resolve().parents[1]
APP = ROOT / "app" / "app.py"


def open_result(path):
    app = AppTest.from_file(str(APP), default_timeout=15).run()
    app.text_input[0].set_value(str(path)).run()
    assert not app.exception
    return app


@pytest.fixture
def toy():
    return json.loads(DEFAULT_PATH.read_text(encoding="utf-8"))


def test_toy_app_and_lookups(toy):
    app = AppTest.from_file(str(APP), default_timeout=15).run()
    assert not app.exception
    assert app.metric[0].value == "OPTIMAL"
    assert len(app.tabs) == 4
    assert set(app.dataframe[1].value["Chỉ số"]) == set(KPI_NOTES)
    assert len(app.dataframe[0].value) == 6
    room = toy["sittings"][0]["rooms"][0]
    app.selectbox[0].select(room["room_id"]).run()
    assert set(app.dataframe[3].value["Phòng"]) == {room["room_id"]}
    student = room["students"][0]
    app.text_input[1].set_value(student).run()
    expected = sum(student in r["students"] for s in toy["sittings"] for r in s["rooms"])
    assert len(app.dataframe[4].value) == expected
    lecturer = room["proctors"][0]
    app.selectbox[1].select(lecturer).run()
    assert all(lecturer in value.split(", ") for value in app.dataframe[5].value["Giám thị"])
    app.text_input[1].set_value("KHONG_TON_TAI").run()
    assert any("Không tìm thấy" in info.value for info in app.info)
    assert not app.exception


def test_upload_takes_priority_over_path(toy):
    toy["status"] = "FEASIBLE"
    toy["kpi"]["avg_gap_days"] = None
    upload = BytesIO(json.dumps(toy).encode("utf-8"))
    with patch("streamlit.file_uploader", return_value=upload):
        app = AppTest.from_file(str(APP), default_timeout=15).run()
        app.text_input[0].set_value("missing.json").run()
    assert not app.exception
    assert app.metric[0].value == "FEASIBLE"
    assert not app.error
    assert "Không áp dụng" in app.dataframe[1].value.loc[2, "Giá trị"]


def test_calendar_has_four_periods_and_repeats_two_period_exam(toy):
    sitting = next(s for s in toy["sittings"] if len(s["periods"]) == 2)
    rows = calendar_rows(toy["sittings"])
    row = next(row for row in rows if row["Ngày"] == sitting["date"])
    assert len(row) == 5
    for period in sitting["periods"]:
        key = next(key for key in row if key.startswith(f"Ca {period} "))
        assert sitting["course_id"] in row[key]
        assert all(r["room_id"] in row[key] for r in sitting["rooms"])


def test_excel_contains_room_assignments_and_students(toy):
    workbook = load_workbook(BytesIO(export_excel(toy["sittings"])))
    assert workbook.sheetnames == ["Lịch tổng", "Sinh viên theo phòng"]
    rooms = [r for s in toy["sittings"] for r in s["rooms"]]
    assert workbook.worksheets[0].max_row == len(rooms) + 1
    assert workbook.worksheets[1].max_row == sum(len(r["students"]) for r in rooms) + 1
    assert workbook.worksheets[1].cell(2, 10).value == rooms[0]["students"][0]
    toy["sittings"][0]["course_name"] = "=1+1"
    workbook = load_workbook(BytesIO(export_excel(toy["sittings"])))
    assert workbook.worksheets[0].cell(2, 3).data_type == "s"


@pytest.mark.parametrize("status", ["INPUT_ERROR", "INFEASIBLE", "UNKNOWN"])
def test_no_schedule_for_unsolved_status(tmp_path, toy, status):
    toy.update(status=status, objective=None, penalties=None, kpi=None)
    # Kể cả file có lượt thi cũ, không vẽ lịch nếu trạng thái không có lời giải.
    toy["messages"] = [{"level": "error", "code": "D01", "message": "Dữ liệu không hợp lệ", "file": "rooms.csv", "line": 2}]
    path = tmp_path / "result.json"
    path.write_text(json.dumps(toy), encoding="utf-8")
    app = open_result(path)
    assert app.metric[0].value == status
    assert app.error
    assert not app.dataframe
    assert not app.tabs
    assert not app.get("download_button")


def test_real_input_error_from_cli(tmp_path):
    data = tmp_path / "bad_data"
    shutil.copytree(ROOT / "data" / "toy", data)
    rooms = data / "rooms.csv"
    rooms.write_text(rooms.read_text(encoding="utf-8").replace("normal,60", "normal,-1"), encoding="utf-8")
    result_path = tmp_path / "loi.json"
    result = subprocess.run(
        [sys.executable, "-m", "src.run", "--data", str(data), "--out", str(result_path)],
        cwd=ROOT, capture_output=True, check=False,
    )
    assert result.returncode == 1, result.stderr
    assert json.loads(result_path.read_text(encoding="utf-8"))["status"] == "INPUT_ERROR"
    app = open_result(result_path)
    assert app.metric[0].value == "INPUT_ERROR"
    assert app.error
    assert not app.dataframe


def test_messages_and_violations(tmp_path, toy):
    toy["messages"] = [{"level": level, "code": level, "message": f"Thông báo {level}"}
                       for level in ("error", "warning", "info")]
    toy["violations"] = [{"code": "A-H2", "message": "Trùng ca", "refs": {"student_id": "22020001"}}]
    path = tmp_path / "violations.json"
    path.write_text(json.dumps(toy), encoding="utf-8")
    app = open_result(path)
    assert any("A-H2" in message.value for message in app.error)
    assert any("warning" in message.value for message in app.warning)
    assert any("info" in message.value for message in app.info)


@pytest.mark.parametrize("content", ["{", "[]", '{"status": "INVALID"}', '{"status": "INPUT_ERROR", "messages": [1]}'])
def test_bad_json_is_reported(tmp_path, content):
    path = tmp_path / "bad.json"
    path.write_text(content, encoding="utf-8")
    assert open_result(path).error


def test_missing_file_is_reported(tmp_path):
    assert open_result(tmp_path / "missing.json").error

